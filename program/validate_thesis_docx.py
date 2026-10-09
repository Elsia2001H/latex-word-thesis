#!/usr/bin/env python3
"""Structural acceptance checks for the generated thesis DOCX."""

from __future__ import annotations

import sys
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from docx import Document


NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def main(path):
    doc = Document(str(path))
    headings = [
        (p.style.name, p.text)
        for p in doc.paragraphs
        if p.style and p.style.name.startswith("Heading")
    ]
    bib_paragraphs = [
        p for p in doc.paragraphs if p.style and p.style.name == "Bibliography"
    ]
    with ZipFile(path) as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
        equations = root.xpath(".//m:oMath", namespaces=NS)
        images = root.xpath(".//a:blip", namespaces=NS)
        sects = root.xpath(".//w:sectPr", namespaces=NS)
        toc_fields = root.xpath(".//w:instrText[contains(., 'TOC')]", namespaces=NS)
        print(f"headings={len(headings)} bibliography={len(bib_paragraphs)} equations={len(equations)} images={len(images)} sections={len(sects)} toc_fields={len(toc_fields)}")
        for i, sect in enumerate(sects, 1):
            fmt = sect.xpath("string(w:pgNumType/@w:fmt)", namespaces=NS) or "inherit"
            start = sect.xpath("string(w:pgNumType/@w:start)", namespaces=NS) or "continue"
            headers = sect.xpath("w:headerReference/@r:id", namespaces={**NS, "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"})
            footers = sect.xpath("w:footerReference/@r:id", namespaces={**NS, "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"})
            print(f"section {i}: page={fmt}/{start} headers={len(headers)} footers={len(footers)}")
        print("heading order:")
        for style, text in headings:
            print(f"  {style}: {text}")

        required = ["摘  要", "ABSTRACT", "目  录", "第一章", "参考文献", "攻读博士学位期间的科研成果", "致  谢"]
        joined = "\n".join(text for _, text in headings)
        missing = [item for item in required if item not in joined]
        if missing or len(bib_paragraphs) < 3 or not equations or not images or len(toc_fields) < 1:
            raise SystemExit(f"validation failed; missing={missing}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: validate_thesis_docx.py FILE.docx")
    main(Path(sys.argv[1]))
