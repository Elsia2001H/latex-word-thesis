"""Regression checks for arbitrary chapter counts and decimal preservation."""
import os
import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from thesis_numbering import chapter_mapping, renumber
from postprocess_thesis_docx import replace_visible_numbers


class NumberingTests(unittest.TestCase):
    def test_four_chapters_two_appendices(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'test.tex'
            source.write_text(r'\chapter*{摘要}' + ''.join(r'\chapter{' + title + '}' for title in ['甲','乙','丙','丁']) + r'\appendix\chapter{补充一}\chapter{补充二}', encoding='utf-8')
            old = os.environ.get('THESIS_TEX')
            os.environ['THESIS_TEX'] = str(source)
            try:
                doc = Document()
                for number, title in enumerate(['甲','乙','丙','丁','补充一','补充二'], 5):
                    doc.add_paragraph('Section %d %s' % (number, title), 'Heading 1')
                self.assertEqual(chapter_mapping(doc), {'5':'1','6':'2','7':'3','8':'4','9':'A','10':'B'})
            finally:
                if old is None: os.environ.pop('THESIS_TEX', None)
                else: os.environ['THESIS_TEX'] = old

    def test_decimal_constants_and_internal_references(self):
        doc = Document()
        paragraph = doc.add_paragraph('常数 5.1 和 8.2 保持原样。')
        link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), 'eq:test')
        run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '(8.2)'
        run.append(text); link.append(run); paragraph._p.append(link)
        replace_visible_numbers(doc, {'5':'1', '8':'4'})
        self.assertEqual(paragraph.runs[0].text, '常数 5.1 和 8.2 保持原样。')
        self.assertEqual(text.text, '(4-2)')
        self.assertEqual(renumber('(10.12)', {'10':'B'}), '(B-12)')


if __name__ == '__main__':
    unittest.main()
