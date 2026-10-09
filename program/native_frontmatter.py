"""Rebuild four fixed-layout front pages as editable Word text boxes/lines."""
import json
from copy import deepcopy
from pathlib import Path
from lxml import etree
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

V = 'urn:schemas-microsoft-com:vml'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def rebuild(document, manifest: Path):
    pages = json.loads(manifest.read_text(encoding='utf-8'))
    abstract = next(p for p in document.paragraphs if ''.join(p.text.split()) == '摘要')
    body = document._element.body
    old = []
    for node in body:
        if node is abstract._p:
            break
        old.append(node)
    breaks = [deepcopy(s) for n in old for s in n.iter(qn('w:sectPr'))]
    if len(breaks) not in (1, 2):
        raise ValueError('Unexpected front-matter section structure')
    for node in old:
        body.remove(node)
    shape_id = 10000
    for index, page in enumerate(pages):
        anchor = document.add_paragraph(style='Normal')
        fmt = anchor.paragraph_format
        fmt.first_line_indent = Pt(0)
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.line_spacing = Pt(1)
        fmt.page_break_before = index > 0 and not (len(breaks) == 2 and index == 3)
        fmt.keep_with_next = False
        anchor.add_run().font.size = Pt(1)
        for span in page['spans']:
            shape_id += 1
            size = span['size']
            font = 'KaiTi' if 'KaiTi' in span['font'] else ('SimHei' if 'SimHei' in span['font'] else ('SimSun' if 'SimSun' in span['font'] else 'Times New Roman'))
            pict = OxmlElement('w:pict')
            shape = etree.SubElement(pict, '{%s}shape' % V)
            shape.set('id', 'native_front_%d' % shape_id)
            shape.set('type', '#_x0000_t202')
            shape.set('stroked', 'f'); shape.set('filled', 'f')
            shape.set('style', 'position:absolute;margin-left:%.3fpt;margin-top:%.3fpt;width:%.3fpt;height:%.3fpt;z-index:1;mso-position-horizontal-relative:page;mso-position-vertical-relative:page' % (span['x'], span['baseline'] - size * 1.045, span['width'] + 2, size * 1.6))
            textbox = etree.SubElement(shape, '{%s}textbox' % V)
            textbox.set('inset', '0,0,0,0')
            content = etree.SubElement(textbox, qn('w:txbxContent'))
            para = etree.SubElement(content, qn('w:p'))
            ppr = etree.SubElement(para, qn('w:pPr'))
            spacing = etree.SubElement(ppr, qn('w:spacing'))
            for name, val in {'before':'0','after':'0','line':str(round(size*1.3*20)),'lineRule':'exact'}.items(): spacing.set(qn('w:'+name), val)
            ind = etree.SubElement(ppr, qn('w:ind')); ind.set(qn('w:firstLine'),'0')
            run = etree.SubElement(para, qn('w:r'))
            rpr = etree.SubElement(run, qn('w:rPr'))
            fonts = etree.SubElement(rpr, qn('w:rFonts'))
            for name in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+name),font)
            sz = etree.SubElement(rpr,qn('w:sz')); sz.set(qn('w:val'),str(round(size*2)))
            bold = etree.SubElement(rpr,qn('w:b')); bold.set(qn('w:val'),'1' if 'Bold' in span['font'] else '0')
            fit = etree.SubElement(rpr,qn('w:fitText')); fit.set(qn('w:val'),str(round(span['width']*20))); fit.set(qn('w:id'),str(shape_id))
            text = etree.SubElement(run, qn('w:t')); text.set(qn('xml:space'),'preserve'); text.text = span['text']
            anchor.add_run()._r.append(pict)
        for x1,y1,x2,y2 in page['lines']:
            pict=OxmlElement('w:pict'); line=etree.SubElement(pict,'{%s}line'%V)
            line.set('style','position:absolute;mso-position-horizontal-relative:page;mso-position-vertical-relative:page')
            line.set('from','%.3fpt,%.3fpt'%(x1,y1)); line.set('to','%.3fpt,%.3fpt'%(x2,y2))
            line.set('strokeweight','0.5pt'); line.set('strokecolor','#000000')
            anchor.add_run()._r.append(pict)
        if len(breaks) == 2 and index == 2:
            anchor._p.get_or_add_pPr().append(breaks[0])
        if index == len(pages)-1:
            anchor._p.get_or_add_pPr().append(breaks[-1])
        abstract._p.addprevious(anchor._p)
    abstract.paragraph_format.page_break_before = False
