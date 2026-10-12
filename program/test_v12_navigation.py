import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from postprocess_thesis_docx import move_bibliography
from finalize_word_layout import format_cached_toc, format_linked_citations

class NavigationTests(unittest.TestCase):
    def test_linked_citation_brackets_and_digits_remain_superscript(self):
        d=Document();p=d.add_paragraph();p.add_run('[')
        link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'ref-test')
        r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);link.append(r);p._p.append(link);p.add_run(']')
        format_linked_citations(p)
        self.assertEqual(p._p.xpath('.//w:vertAlign/@w:val'),['superscript']*3)

    def test_bibliography_bookmark_moves_with_entry(self):
        d=Document()
        if 'Bibliography' not in d.styles:d.styles.add_style('Bibliography',WD_STYLE_TYPE.PARAGRAPH)
        target=d.add_paragraph('攻读博士学位期间的科研成果',style='Heading 1')
        p=d.add_paragraph('[1] Example',style='Bibliography')
        a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),'50');a.set(qn('w:name'),'ref-example');p._p.addprevious(a)
        b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),'50');p._p.addnext(b)
        move_bibliography(d)
        self.assertIs(p._p.getnext(),target._p)
        self.assertEqual(p._p.xpath('./w:bookmarkStart/@w:name'),['ref-example'])
        self.assertEqual(p._p.xpath('./w:bookmarkEnd/@w:id'),['50'])

    def test_toc_title_black_font_page_regular(self):
        d=Document();s=d.styles.add_style('TOC 1',WD_STYLE_TYPE.PARAGRAPH);s.style_id='TOC1'
        p=d.add_paragraph(style=s);p.add_run('第一章 测试');p.add_run().add_tab();p.add_run('1')
        with TemporaryDirectory() as folder:
            path=Path(folder)/'toc.docx';d.save(path);format_cached_toc(path);d=Document(path)
        r=d.paragraphs[0]._p.xpath('./w:r')
        self.assertEqual(r[0].xpath('./w:rPr/w:rFonts/@w:eastAsia'),['SimHei'])
        self.assertEqual(r[0].xpath('./w:rPr/w:b/@w:val'),['1'])
        self.assertEqual(r[-1].xpath('./w:rPr/w:b/@w:val'),['0'])

if __name__=='__main__':unittest.main()
