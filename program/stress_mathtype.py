"""Generate a repeatable large formula fixture without changing thesis templates."""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from mathtype_convert import convert

CASES=[r'\frac{a+b}{c+d}',r'\int_0^T e^{-\lambda t}q(t)^2\,\mathrm{d}t',
       r'\begin{bmatrix}0&1&0\\1&0&1\\0&1&0\end{bmatrix}',
       r'\sqrt{\sum_{i=1}^{N}(x_i-\bar{x})^2}',
       r'\begin{aligned}a&=b+c,\\d&=e+f,\\g&=h+i.\end{aligned}',
       r'f(x)=\begin{cases}0,&x<0,\\x^2,&x\geq0.\end{cases}',
       r'\frac{\sum_{i=1}^{N}w_i(x_i-\bar{x})^2+\sum_{j=1}^{M}v_j(y_j-\bar{y})^2}{\alpha+\beta+\sqrt{\sum_{i=1}^{N}w_i^2}}']


def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--count',type=int,default=513)
    args=parser.parse_args(); root=Path(__file__).parents[1]/'qa/stress'; root.mkdir(parents=True,exist_ok=True)
    doc=Document(); manifest=[]
    for i in range(args.count):
        tex=CASES[i%len(CASES)]; mode=i%4
        if mode==0:
            paragraph=doc.add_paragraph('Inline %d '%i)
        elif mode==1:
            paragraph=doc.add_table(rows=1,cols=1).cell(0,0).paragraphs[0]
        elif mode==2:
            paragraph=doc.add_paragraph('Caption %d '%i,style='Caption')
        else:
            paragraph=doc.add_paragraph()
        math=OxmlElement('m:oMath'); run=OxmlElement('m:r'); text=OxmlElement('m:t')
        text.text=tex; run.append(text); math.append(run); paragraph._p.append(math)
        if mode==3:
            paragraph.add_run(' (%d-%d)'%(i//100+1,i%100+1))
        manifest.append({'tex':tex,'display':mode==3})
    output=root/('stress-%d.docx'%args.count); source=root/('stress-%d.json'%args.count)
    doc.save(output); source.write_text(json.dumps(manifest),encoding='utf-8')
    convert(output,source)


if __name__=='__main__':
    main()
