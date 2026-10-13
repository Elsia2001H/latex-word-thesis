"""Replace OMML slots only, preserving the existing thesis layout and numbering."""
import argparse
import copy
import ctypes as C
import json
import os
import posixpath
import re
import struct
import subprocess
import time
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
from docx import Document
from mathtype_cfb import read_streams, replace_native
from mathtype_alignment import align_native
from mathtype_discovery import discover

NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm':'http://schemas.openxmlformats.org/officeDocument/2006/math',
      'o':'urn:schemas-microsoft-com:office:office', 'v':'urn:schemas-microsoft-com:vml',
      'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
CORE = Path(__file__).resolve().parent


def progress(stage, done, total):
    print('[MATHTYPE_PROGRESS]'+json.dumps({'stage':stage,'done':done,'total':total}), flush=True)


def normalize(tex):
    tex = re.sub(r'\{\\qquad\([^()]+\)\}\s*$', '', tex)
    tex = re.sub(r'\\(?:begin|end)\{equation\*?\}', '', tex)
    tex = re.sub(r'\\(?:label|tag)\{[^{}]*\}', '', tex)
    tex = re.sub(r'\\(?:lvert|rvert)\b', '|', tex)
    tex = re.sub(r'\\(?:lVert|rVert)\b', lambda m:r'\|', tex)
    return tex.replace(r'\bm{',r'\mathbf{').replace(r'\begin{aligned}',r'\begin{array}{rl}').replace(r'\end{aligned}',r'\end{array}').strip()


def preferences(size):
    """Write our own settings, using documented MathType INI keys."""
    return '''[Styles]
Text=Times New Roman
Function=Times New Roman
Variable=Times New Roman,I
LCGreek=Symbol,I
UCGreek=Symbol
Symbol=Symbol
Vector=Times New Roman,B
Number=Times New Roman
User1=Courier New
User2=Times New Roman
MTExtra=MT Extra
[Sizes]
Full=%s pt
Script=58 %%
ScriptScript=42 %%
Symbol=150 %%
SubSymbol=100 %%
User1=75 %%
User2=150 %%
SmallLargeIncr=1 pt
[Spacing]
LineSpacing=150 %%
MatrixRowSpacing=150 %%
MatrixColSpacing=100 %%
SuperscriptHeight=45 %%
SubscriptDepth=25 %%
SubSupGap=8 %%
LimHeight=25 %%
LimDepth=100 %%
LimLineSpacing=100 %%
NumerHeight=35 %%
DenomDepth=100 %%
FractBarOver=8 %%
FractBarThick=5 %%
SubFractBarThick=2.5 %%
FractGap=8 %%
FenceOver=8 %%
OperSpacing=100 %%
NonOperSpacing=100 %%
CharWidth=0 %%
MinGap=8 %%
VertRadGap=17 %%
HorizRadGap=8 %%
RadWidth=100 %%
EmbellGap=12.5 %%
PrimeHeight=45 %%
BoxStrokeThick=5 %%
StikeThruThick=5 %%
MatrixLineThick=5 %%
RadStrokeThick=5 %%
HorizFenceGap=10 %%
''' % size


def wmf_math(wmf):
    if len(wmf)<40 or struct.unpack_from('<I',wmf)[0] != 0x9ac6cdd7:
        raise ValueError('Invalid placeable equation WMF')
    unsupported = re.search(rb'<(?:\w+:)?mtext>[^<]*(?:&#x005[Cc];|&#92;|\\)[^<]*</(?:\w+:)?mtext>', wmf)
    if unsupported:
        raise ValueError('MathType left an unsupported TeX command as text: '+unsupported.group().decode('ascii',errors='replace'))
    chunks, total, pos = [], None, 40
    while pos+6 <= len(wmf):
        words, func = struct.unpack_from('<IH',wmf,pos)
        if words<3 or pos+words*2>len(wmf):
            raise ValueError('Invalid WMF record')
        if func==0x626:
            payload = wmf[pos+10:pos+words*2]
            if payload.startswith(b'AppsMFCC') and len(payload)>=19 and payload[8:10]==b'\x01\x00':
                expected, part = struct.unpack_from('<II',payload,10)
                start = payload.index(b'\0',18)+1
                if total is None and not payload[start:start+2]==b'\x05\x01':
                    pos += words*2
                    continue
                if total is not None and total != expected:
                    raise ValueError('Inconsistent MTEF chunks')
                if len(payload)-start < part:
                    raise ValueError('Truncated MTEF chunk')
                total = expected; chunks.append(payload[start:start+part])
                if sum(map(len,chunks))==total:
                    break
        pos += words*2
        if func==0:
            break
    mtef = b''.join(chunks)
    if not mtef or len(mtef)!=total:
        raise ValueError('Incomplete MTEF payload')
    h=struct.unpack('<IHhhhhHIH',wmf[:22])
    return mtef, (h[4]-h[2])*72/h[6], (h[5]-h[3])*72/h[6]


class Formatter:
    class Dims(C.Structure):
        _fields_=[('baseline',C.c_short),('left',C.c_int),('top',C.c_int),('right',C.c_int),('bottom',C.c_int)]
    def __init__(self, folder):
        dll_path=discover().require()
        self.dll=C.WinDLL(str(dll_path)); d=self.dll
        d.MTInitAPI.argtypes=[C.c_short,C.c_short]
        d.MTGetPrefsFromFile.argtypes=[C.c_char_p,C.c_void_p,C.c_short]
        d.MTXFormSetPrefs.argtypes=[C.c_short,C.c_char_p]
        d.MTXFormEqn.argtypes=[C.c_short,C.c_short,C.c_void_p,C.c_int,C.c_short,C.c_short,C.c_void_p,C.c_int,C.c_char_p,C.POINTER(self.Dims)]
        if d.MTInitAPI(0,30)!=1:
            raise RuntimeError('MathType API initialization failed')
        self.prefs={}; self.current=None
        for name,size in [('body',12),('small',10.5)]:
            path=folder/(name+'.eqp')
            retained=CORE.parent/'formula-styles'/(name+'.eqp')
            path.write_bytes(retained.read_bytes() if retained.exists() else preferences(size).encode('ascii'))
            buf=C.create_string_buffer(8192)
            if d.MTGetPrefsFromFile(str(path).encode('mbcs'),buf,8192)!=0:
                self.close(); raise RuntimeError('Cannot load equation profile '+name)
            self.prefs[name]=buf.value
    def close(self):
        self.dll.MTTermAPI()
    def format(self, native, profile, path):
        d=self.dll
        if self.current!=profile:
            if d.MTXFormReset()!=0 or d.MTXFormSetPrefs(3,self.prefs[profile])!=0:
                raise RuntimeError('Cannot apply MathType profile')
            self.current=profile
        length=struct.unpack_from('<I',native,8)[0]
        if not 0<length<=len(native)-28:
            raise ValueError('Invalid native MTEF length')
        mtef=native[28:28+length]; buf=C.create_string_buffer(mtef)
        dummy=C.create_string_buffer(1); dims=self.Dims()
        if d.MTXFormEqn(-3,4,buf,len(mtef),-4,6,dummy,1,str(path).encode('mbcs'),C.byref(dims))!=0:
            raise RuntimeError('MathType formatting failed')
        return path.read_bytes()


def files_of(path):
    with ZipFile(path) as z:
        return {n:z.read(n) for n in z.namelist()}


def convert(path, manifest, batch_size=100):
    # MathType and Word automation use shared application state. Serialize even
    # direct CLI tests, in addition to the launcher-wide conversion mutex.
    kernel=C.WinDLL('kernel32',use_last_error=True)
    kernel.CreateMutexW.argtypes=[C.c_void_p,C.c_bool,C.c_wchar_p]; kernel.CreateMutexW.restype=C.c_void_p
    kernel.CloseHandle.argtypes=[C.c_void_p]; kernel.ReleaseMutex.argtypes=[C.c_void_p]
    handle=kernel.CreateMutexW(None,True,'Local\\ThesisMathTypeFormulaWorker')
    if not handle:
        raise C.WinError(C.get_last_error())
    if C.get_last_error()==183:
        kernel.CloseHandle(handle)
        raise RuntimeError('Another MathType conversion is running; retry when it finishes')
    try:
        return _convert(path,manifest,batch_size)
    finally:
        kernel.ReleaseMutex(handle); kernel.CloseHandle(handle)


def _convert(path, manifest, batch_size):
    started=time.perf_counter(); files=files_of(path)
    xml=E.fromstring(files['word/document.xml'])
    slots=xml.xpath('//m:oMath',namespaces=NS)
    section=Document(path).sections[0]
    body_width=(section.page_width-section.left_margin-section.right_margin)/12700
    body_height=(section.page_height-section.top_margin-section.bottom_margin)/12700
    formulas=json.loads(Path(manifest).read_text(encoding='utf-8'))
    if len(slots)!=len(formulas):
        raise ValueError('Math manifest/OMML count mismatch: %d/%d' % (len(formulas),len(slots)))
    if not slots:
        progress('complete',0,0); return
    if not 1<=batch_size<=100:
        raise ValueError('Batch size must be between 1 and 100')
    work=path.parent/(path.stem+'.mathtype-work-'+str(time.time_ns())); work.mkdir(exist_ok=False)
    rels=E.fromstring(files['word/_rels/document.xml.rels'])
    content=E.fromstring(files['[Content_Types].xml'])
    formatter=Formatter(work); audit=[]; objects=[]
    try:
        for base in range(0,len(formulas),batch_size):
            batch=formulas[base:base+batch_size]; batch_path=work/('batch-%04d.docx'%base)
            doc=Document()
            for case in batch:
                tex=normalize(case['tex'])
                if not tex:
                    raise ValueError('Empty equation')
                doc.add_paragraph(('\\['+tex+'\\]') if case['display'] else ('$'+tex+'$'))
            doc.save(batch_path); progress('convert',base,len(formulas))
            owner=work/('batch-%04d-owner.json'%base)
            try:
                subprocess.run(['powershell','-NoProfile','-ExecutionPolicy','Bypass','-File',str(CORE/'convert_math_batch.ps1'),'-Path',str(batch_path),'-Expected',str(len(batch)),'-OwnerFile',str(owner)],check=True,timeout=180)
            except subprocess.TimeoutExpired:
                # Stop only the exact background process created by this batch;
                # verify its PID, creation time and executable before cleanup.
                cleanup='''param([string]$OwnerFile)
$info=Get-Content -LiteralPath $OwnerFile -Raw | ConvertFrom-Json
$p=Get-Process -Id $info.pid -ErrorAction SilentlyContinue
if($p -and $p.ProcessName -eq 'WINWORD' -and $p.StartTime.ToUniversalTime().Ticks -eq $info.startTicks){Stop-Process -Id $p.Id -Force}
'''
                script=work/'cleanup-owned-word.ps1'; script.write_text(cleanup,encoding='utf-8')
                if owner.exists():
                    subprocess.run(['powershell','-NoProfile','-ExecutionPolicy','Bypass','-File',str(script),'-OwnerFile',str(owner)],check=False,timeout=20)
                raise RuntimeError('MathType batch timed out; native document retained, batch diagnostics: '+str(work)) from None
            converted=files_of(batch_path); bx=E.fromstring(converted['word/document.xml'])
            br={r.get('Id'):r.get('Target') for r in E.fromstring(converted['word/_rels/document.xml.rels'])}
            runs=bx.xpath('//w:r[w:object]',namespaces=NS)
            if len(runs)!=len(batch):
                raise ValueError('Converted equation count mismatch')
            for offset,run in enumerate(runs):
                i=base+offset; slot=slots[i]
                small=bool(slot.xpath('ancestor::w:tc',namespaces=NS))
                paragraph=slot.xpath('ancestor::w:p',namespaces=NS)[-1]
                styles=paragraph.xpath('./w:pPr/w:pStyle/@w:val',namespaces=NS)
                small=small or any('caption' in s.lower() for s in styles)
                profile='small' if small else 'body'
                run=copy.deepcopy(run); obj=run.find('w:object',NS)
                ole=obj.find('o:OLEObject',NS); shape=obj.find('v:shape',NS)
                preview=shape.find('v:imagedata',NS)
                binary=converted[posixpath.normpath('word/'+br[ole.get('{'+NS['r']+'}id')])]
                streams,_=read_streams(binary)
                native=align_native(streams['Equation Native'],formulas[i]['tex'])
                wmf=formatter.format(native,profile,work/('equation-%04d.wmf'%i))
                mtef,width,height=wmf_math(wmf)
                available=body_width-48 if formulas[i]['display'] else body_width
                cell=slot.xpath('ancestor::w:tc',namespaces=NS)
                if cell:
                    cell_width=cell[-1].find('w:tcPr/w:tcW',NS)
                    if cell_width is not None and cell_width.get('{'+NS['w']+'}type')=='dxa':
                        available=min(available,int(cell_width.get('{'+NS['w']+'}w'))/20-12)
                if width>available or height>body_height-24:
                    raise ValueError('Equation %d exceeds its layout area (%.1f x %.1f pt); split it explicitly with aligned/array in TeX. Native document retained.'%(i+1,width,height))
                binary=replace_native(binary,mtef)
                for element, ext, payload, kind in [(ole,'bin',binary,'oleObject'),(preview,'wmf',wmf,'image')]:
                    target=('embeddings/' if ext=='bin' else 'media/')+'mt-equation-%04d.'%i+ext
                    rid='rIdMT%d%s'%(i,ext)
                    if any(r.get('Id')==rid for r in rels):
                        raise ValueError('Equation relationship collision')
                    E.SubElement(rels,'{'+REL+'}Relationship',Id=rid,Type=NS['r']+'/'+kind,Target=target)
                    element.set('{'+NS['r']+'}id',rid); files['word/'+target]=payload
                shape_id='_x0000_mtshape%d'%i
                shape.set('id',shape_id); ole.set('ShapeID',shape_id); ole.set('ObjectID','_%d'%(1000000+i))
                style=re.sub(r'width:[\d.]+pt','width:%.4fpt'%width,shape.get('style'))
                style=re.sub(r'height:[\d.]+pt','height:%.4fpt'%height,style); shape.set('style',style)
                if small:
                    position=run.find('w:rPr/w:position',NS)
                    if position is not None:
                        position.set('{'+NS['w']+'}val',str(round(int(position.get('{'+NS['w']+'}val'))*10.5/12)))
                objects.append(run); audit.append({'index':i,'profile':profile,'width':width,'height':height,'mtef_bytes':len(mtef)})
                progress('format',i+1,len(formulas))
        # Only replace mathematical nodes. Original paragraphs, table geometry,
        # bookmark numbers, hyperlinks and other package parts remain unchanged.
        for slot,run in zip(slots,objects):
            slot.getparent().replace(slot,run)
        for para in xml.xpath('//m:oMathPara',namespaces=NS):
            parent=para.getparent(); index=parent.index(para)
            for run in list(para):
                if run.tag=='{'+NS['w']+'}r':
                    parent.insert(index,run); index+=1
            parent.remove(para)
        if xml.xpath('//m:oMath',namespaces=NS):
            raise ValueError('Unconverted math remains')
        for ext,ctype in [('bin','application/vnd.openxmlformats-officedocument.oleObject'),('wmf','image/x-wmf')]:
            if not any(e.get('Extension')==ext for e in content):
                E.SubElement(content,'{'+content.nsmap[None]+'}Default',Extension=ext,ContentType=ctype)
        for name,tree in [('word/document.xml',xml),('word/_rels/document.xml.rels',rels),('[Content_Types].xml',content)]:
            files[name]=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
        candidate=work/'converted.docx'
        with ZipFile(candidate,'w',ZIP_DEFLATED) as z:
            for name,data in files.items():
                z.writestr(name,data)
        if len(E.fromstring(files_of(candidate)['word/document.xml']).xpath('//o:OLEObject[@ProgID="Equation.DSMT4"]',namespaces=NS))!=len(formulas):
            raise ValueError('Final MathType object count mismatch')
        report={'equations':audit,'seconds':time.perf_counter()-started,'batch_size':batch_size}
        (work/'audit.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
        path.with_suffix('.native.docx').write_bytes(path.read_bytes())
        os.replace(candidate,path)
        progress('complete',len(formulas),len(formulas))
    finally:
        formatter.close()


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('docx',type=Path); parser.add_argument('manifest',type=Path)
    parser.add_argument('--batch-size',type=int,default=100)
    args=parser.parse_args(); convert(args.docx.resolve(),args.manifest.resolve(),args.batch_size)
